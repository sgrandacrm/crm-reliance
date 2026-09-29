# -*- coding: utf-8 -*-
"""
PROCESO 1 - OPERACIONES
Genera los archivos DBroker desde los archivos CASH del dia:
UNIFICADO, CONTRATANTES, POLIZAS VH/IN/AP/VIDA, MATRICES VH/IN/AP/VIDA, LOG
"""
import sys
import io
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from unificar import unificar
from gen_contratantes import (
    cargar_enriquecimiento, generar_contratantes,
    gen_norm, fecha_ddmmyyyy, celu_norm,
)
from gen_polizas_vh import generar_polizas_vh
from gen_matrices_vh import generar_matrices_vh
from gen_polizas_in import generar_polizas_in
from gen_polizas_ap_vida import generar_polizas_ap_vida
from gen_matrices_in import generar_matrices_in
from gen_matrices_ap_vida import generar_matrices_ap_vida

BASE          = Path(__file__).parent.parent
DIR_ENTRADA   = BASE / "OPERACIONES" / "entrada"
DIR_CATALOGOS = BASE / "catalogos"
DIR_BD_ENRIQ  = BASE / "bd_enriquecimiento"
DIR_SALIDA    = BASE / "OPERACIONES" / "salida"
CATALOGO_PATH = DIR_CATALOGOS / "CATALOGO CONTRATANTES.xls"

COLOR_HEADER = "20584E"
FILL_BLOQ    = PatternFill("solid", fgColor="FFCCCC")
FILL_REV     = PatternFill("solid", fgColor="FFE5B4")
FILL_ALERT   = PatternFill("solid", fgColor="FFF2CC")

COLS_META = [
    "ARCHIVO_ORIGEN", "HOJA_ORIGEN", "FILA_ORIGEN", "LAYOUT",
    "POLIZA_EMITIDA_ORIGINAL", "FACTURA_ASEGURADORA_ORIGINAL",
    "FECHA_NACIMIENTO_NORM", "EDAD_CALCULADA",
    "CLIENTE_MULTIPOLIZA", "ALERTAS", "NUM_ALERTAS", "CAMPOS_BD",
]

COLS_COMUNES = [
    "EJECUTIVO", "CI", "GARANTIA", "NOMBRE_CLIENTE", "TIPO", "CIUDAD",
    "OBS_POLIZA", "CELULAR_1", "CUENTA", "COMENTARIO", "ESTADO_RENOVACION",
    "POLIZA_EMITIDA", "ASEGURADORA_EMISORA", "RAMO",
    "VIGENCIA_DESDE", "VIGENCIA_HASTA",
    "CUOTA_INICIAL", "FACTURA_ASEGURADORA", "NUM_CUOTAS",
    "FECHA_COBRO_CI", "FECHA_COBRO_CF",
    "VALOR_ASEGURADO", "TASA", "PRIMA_NETA", "COMISION",
    "DERECHOS_EMISION", "SUPER_BANCOS", "SEGURO_CAMPESINO", "IVA", "PRIMA_TOTAL",
    "POLIZA_ANTERIOR", "ASEGURADORA_ANTERIOR",
    "CUENTA_BANCO", "NUM_PRESTAMO", "ESTADO_CREDITO",
    "FECHA_DESEMBOLSO", "MONTO_CREDITO", "SALDO_CREDITO", "FECHA_VENC_CREDITO",
    "DIRECCION_DOM", "TELF_DOM", "DIRECCION_OFI", "TELF_OFI",
    "CELULAR_BANCO", "CORREO", "NOMBRE_OFICIAL", "ASEGURADORA_BANCO",
    "TELEFONOS_ADICIONALES", "FECHA_NACIMIENTO", "EDAD_ORIGINAL",
]

COLS_VH_ANALISIS = [
    "REGION", "AXA", "VIDA_ADICIONAL",
    "ULTIMO_VALOR_ASEG", "VALOR_DEPRECIADO",
    "TASA_PRESUPUESTO", "PN_PRESUPUESTO", "COM_PRESUPUESTO",
    "TASA_VIG_ANTERIOR", "TASA_RENOV_ASEG_ACTUAL",
    "CELULAR_2", "GENERO", "ESTADO_CIVIL", "PROFESION",
]

COLS_VH_VEHICULO = [
    "MARCA", "MODELO", "ANIO_VEHICULO", "MOTOR", "CHASIS", "COLOR", "PLACA",
]

COLS_IN_DATOS = [
    "DIRECCION_INMUEBLE", "CONTENIDO", "ROBO",
    "SUBTOTAL_IN", "CUOTAS_IN", "VALOR_CUOTA_IN", "SALDO_IN",
    "COMISION_DESGRAVAMEN", "TASA_DESGRAVAMEN", "PRIMA_NETA_DESGRAVAMEN",
    "SBS_DESGRAVAMEN", "SC_DESGRAVAMEN", "DE_DESGRAVAMEN", "TOTAL_DESGRAVAMEN",
    "CUOTAS_DESGRAVAMEN", "VALOR_CUOTA_DESGRAVAMEN",
    "TOTAL_PRIMAS_A_PAGAR", "TOTAL_DEBITO_MENSUAL", "CARTERA",
]

RAMOS_VH = {"VEHICULOS", "AP", "VIDA"}

COLS_FINANCIERAS_UNIF = [
    "VALOR_ASEGURADO", "TASA", "PRIMA_NETA", "COMISION",
    "DERECHOS_EMISION", "SUPER_BANCOS", "SEGURO_CAMPESINO", "IVA", "PRIMA_TOTAL",
    "AXA", "VIDA_ADICIONAL", "CUOTA_INICIAL",
    "ULTIMO_VALOR_ASEG", "VALOR_DEPRECIADO",
    "TASA_PRESUPUESTO", "PN_PRESUPUESTO", "COM_PRESUPUESTO",
    "TASA_VIG_ANTERIOR", "TASA_RENOV_ASEG_ACTUAL",
    "MONTO_CREDITO", "SALDO_CREDITO",
]

COLS_FINANCIERAS_PVEH = [
    "VAL_ASEGURADO", "TASA", "PRIMA_NETA",
    "DERECHOS_EMISION", "SUPER_BANCOS", "SEGURO_CAMPESINO",
    "VALOR_IVA", "PRIMA_TOTAL", "AXA", "OTRO_VALOR",
]

COLS_FINANCIERAS_PIN = [
    "VAL_ASEGURADO", "TASA", "PRIMA_NETA",
    "DERECHOS_EMISION", "SUPER_BANCOS", "SEGURO_CAMPESINO",
    "VALOR_IVA", "PRIMA_TOTAL",
]

_ALERTAS_POR_CAMPO = {
    "GENERO":       ["GENERO no disponible"],
    "ESTADO_CIVIL": ["ESTADO_CIVIL no disponible"],
    "FECHA_NAC":    ["FECHA_NAC vacia", "FECHA_NAC no encontrada en fuente",
                     "FECHA_NAC serial invalido", "FECHA_NAC formato no reconocido"],
    "DIRECCION":    ["DIRECCION_DOM vacia"],
}


def _cols_para_ramo(ramo: str, disponibles: list) -> list:
    ramo_up = ramo.strip().upper()
    if ramo_up == "VEHICULOS":
        wanted = COLS_COMUNES + COLS_VH_ANALISIS + COLS_VH_VEHICULO + COLS_META
    elif ramo_up in RAMOS_VH:
        wanted = COLS_COMUNES + COLS_VH_ANALISIS + COLS_META
    else:
        wanted = COLS_COMUNES + COLS_IN_DATOS + COLS_META
    return [c for c in wanted if c in disponibles]


def _aplicar_formato(wb):
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF", size=8)
            c.fill = PatternFill("solid", fgColor=COLOR_HEADER)
        ws.freeze_panes = "E2" if ws.title == "MAESTRO" else "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=8)
                if isinstance(c.value, (datetime, date)):
                    c.number_format = "DD/MM/YYYY"


def _colorear_alertas(ws, col_num="NUM_ALERTAS", col_nivel="NIVEL"):
    headers = [str(c.value or "") for c in ws[1]]
    if col_nivel in headers:
        idx = headers.index(col_nivel)
        for row in ws.iter_rows(min_row=2):
            nivel = str(row[idx].value or "").strip().upper()
            fill = (FILL_BLOQ if nivel == "BLOQUEANTE" else
                    FILL_REV  if nivel == "REVISAR"    else
                    FILL_ALERT)
            for cell in row:
                cell.fill = fill
    elif col_num in headers:
        idx = headers.index(col_num)
        for row in ws.iter_rows(min_row=2):
            try:
                n = int(row[idx].value or 0)
            except (ValueError, TypeError):
                n = 0
            if n > 0:
                for cell in row:
                    cell.fill = FILL_ALERT
    else:
        for row in ws.iter_rows(min_row=2):
            if any(cell.value is not None for cell in row):
                for cell in row:
                    cell.fill = FILL_ALERT


def _formato_decimal_cols(ws, nombres):
    cabecera = [c.value for c in ws[1]]
    for nombre in nombres:
        if nombre in cabecera:
            idx = cabecera.index(nombre) + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=idx)
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "0.00"


def _formato_texto_cols(ws, nombres):
    cabecera = [c.value for c in ws[1]]
    for nombre in nombres:
        if nombre in cabecera:
            idx = cabecera.index(nombre) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).number_format = "@"


def _limpiar_alertas_resueltas(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for i, row in df.iterrows():
        campos = str(row.get("CAMPOS_BD") or "")
        if not campos:
            continue
        alertas_raw = str(row.get("ALERTAS") or "")
        if not alertas_raw:
            continue
        partes = [a.strip() for a in alertas_raw.split("|") if a.strip()]
        for campo, fragmentos in _ALERTAS_POR_CAMPO.items():
            if campo in campos:
                partes = [
                    p for p in partes
                    if not any(p.startswith(f) for f in fragmentos)
                ]
        df.at[i, "ALERTAS"]     = " | ".join(partes) if partes else None
        df.at[i, "NUM_ALERTAS"] = len(partes)
    return df


def _vacio(v):
    if v is None:
        return True
    try:
        if pd.isna(v):
            return True
    except (TypeError, ValueError):
        pass
    return str(v).strip().lower() in ("", "nan", "none")


def enriquecer_maestro(df: pd.DataFrame, enr_d: dict, hoy: date) -> pd.DataFrame:
    df = df.copy()
    campos_bd_list = []

    for i, row in df.iterrows():
        ci = str(row["CI"]).strip() if row["CI"] else ""
        enr = enr_d.get(ci)
        campos = []

        if not enr:
            campos_bd_list.append(None)
            continue

        def _aplicar(col_df, bd_val, etiqueta):
            if not bd_val:
                return
            actual = str(row.get(col_df) or "").strip()
            bd_str = str(bd_val).strip()
            df.at[i, col_df] = bd_val
            if actual != bd_str and actual.lower() not in ("", "nan", "none", "no encontrado"):
                campos.append(f"{etiqueta}*")
            elif actual.lower() in ("", "nan", "none", "no encontrado") or not actual:
                campos.append(etiqueta)

        if enr["genero"]:
            nuevo = "MASCULINO" if enr["genero"] == "M" else "FEMENINO"
            actual_g = str(row.get("GENERO") or "").strip()
            df.at[i, "GENERO"] = nuevo
            if actual_g != nuevo:
                etiq = "GENERO*" if actual_g and actual_g.upper() not in ("NO ENCONTRADO", "NAN", "NONE", "") else "GENERO"
                campos.append(etiq)

        if enr["fnac"]:
            df.at[i, "FECHA_NACIMIENTO_NORM"] = enr["fnac"]
            fn = enr["fnac"]
            df.at[i, "EDAD_CALCULADA"] = (
                hoy.year - fn.year - ((hoy.month, hoy.day) < (fn.month, fn.day))
            )
            if _vacio(row["FECHA_NACIMIENTO_NORM"]):
                campos.append("FECHA_NAC")

        _aplicar("ESTADO_CIVIL",  enr["estado_civil"], "ESTADO_CIVIL")
        _aplicar("PROFESION",     enr["profesion"],    "PROFESION")
        _aplicar("DIRECCION_DOM", enr["direccion"],    "DIRECCION")
        _aplicar("CORREO",        enr["mail"],         "CORREO")
        if _vacio(row["CELULAR_1"]) and enr["celu"]:
            df.at[i, "CELULAR_1"] = enr["celu"]
            campos.append("CELULAR")

        campos_bd_list.append(", ".join(campos) if campos else None)

    df["CAMPOS_BD"] = campos_bd_list
    return df


def _cols_alertas(df: pd.DataFrame) -> pd.DataFrame:
    cols = ["ARCHIVO_ORIGEN", "HOJA_ORIGEN", "FILA_ORIGEN", "EJECUTIVO",
            "CI", "NOMBRE_CLIENTE", "RAMO", "POLIZA_EMITIDA", "ALERTAS", "CAMPOS_BD"]
    return df[df["NUM_ALERTAS"] > 0][[c for c in cols if c in df.columns]]


def _guardar_vh(df: pd.DataFrame, ruta: Path):
    cols = _cols_para_ramo("VEHICULOS", df.columns.tolist())
    pivot = df.pivot_table(index="EJECUTIVO", values="CI", aggfunc="count")
    pivot.columns = ["TOTAL"]
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        df[cols].to_excel(xw, sheet_name="VEHICULOS", index=False)
        pivot.reset_index().to_excel(xw, sheet_name="RESUMEN", index=False)
        _cols_alertas(df).to_excel(xw, sheet_name="ALERTAS", index=False)
    wb = load_workbook(ruta)
    _aplicar_formato(wb)
    _formato_decimal_cols(wb["VEHICULOS"], COLS_FINANCIERAS_UNIF)
    _colorear_alertas(wb["VEHICULOS"])
    _colorear_alertas(wb["ALERTAS"])
    wb.save(ruta)


def _guardar_otros(df: pd.DataFrame, ruta: Path):
    ramos = df["RAMO"].dropna().unique()
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        for ramo in sorted(ramos):
            df_r = df[df["RAMO"] == ramo]
            cols = _cols_para_ramo(ramo, df.columns.tolist())
            nombre_hoja = str(ramo)[:31]
            df_r[cols].to_excel(xw, sheet_name=nombre_hoja, index=False)
        _cols_alertas(df).to_excel(xw, sheet_name="ALERTAS", index=False)
    wb = load_workbook(ruta)
    _aplicar_formato(wb)
    for ws in wb.worksheets:
        if ws.title != "ALERTAS":
            _formato_decimal_cols(ws, COLS_FINANCIERAS_UNIF)
            _colorear_alertas(ws)
    _colorear_alertas(wb["ALERTAS"])
    wb.save(ruta)


def run():
    hoy       = date.today()
    fecha_str = hoy.strftime("%Y-%m-%d")

    print("=" * 60)
    print(f"OPERACIONES DBroker  |  {hoy.strftime('%d/%m/%Y')}")
    print("=" * 60)

    # [1/11] Verificar archivos de entrada
    archivos_cash = sorted(p for p in DIR_ENTRADA.glob("*.xls*")
                           if p.suffix.lower() in (".xlsx", ".xlsm"))
    if not archivos_cash:
        print("\nERROR: no se encontraron archivos .xlsx en la carpeta 'entrada'.")
        print("       Coloca los archivos CASH de los asesores ahi y vuelve a ejecutar.")
        return

    print(f"\n[1/11] Archivos CASH en entrada: {len(archivos_cash)}")
    for f in archivos_cash:
        print(f"       - {f.name}")

    if not CATALOGO_PATH.exists():
        print(f"\nERROR: no se encontro el catalogo en:\n  {CATALOGO_PATH}")
        return

    bd_archivos = sorted(DIR_BD_ENRIQ.glob("*.xlsx"))
    if not bd_archivos:
        print("\nERROR: no se encontro la BD de enriquecimiento en 'bd_enriquecimiento'.")
        return
    enriq_path = bd_archivos[0]

    # [2/11] Cargar BD de enriquecimiento
    print(f"\n[2/11] Cargando BD de enriquecimiento: {enriq_path.name}")
    enr_d = cargar_enriquecimiento(enriq_path)
    print(f"       {len(enr_d):,} cedulas cargadas")

    # [3/11] Unificar archivos CASH
    print("\n[3/11] Unificando archivos CASH...")
    df_maestro, log_unif = unificar(archivos_cash, hoy=hoy)
    for linea in log_unif:
        print(f"       {linea}")
    if df_maestro.empty:
        print("\nERROR: no se pudo extraer ninguna fila. Revisa los archivos de entrada.")
        return
    print(f"\n       Filas unificadas   : {len(df_maestro)}")
    print(f"       Clientes unicos CI : {df_maestro['CI'].nunique()}")

    # [4/11] Enriquecer MAESTRO desde BD
    print("\n[4/11] Enriqueciendo MAESTRO desde BD...")
    df_maestro = enriquecer_maestro(df_maestro, enr_d, hoy)
    df_maestro = _limpiar_alertas_resueltas(df_maestro)
    enriq_filas = df_maestro["CAMPOS_BD"].notna().sum()
    campos_cnt = {}
    for v in df_maestro["CAMPOS_BD"].dropna():
        for c in v.split(", "):
            campos_cnt[c] = campos_cnt.get(c, 0) + 1
    print(f"       Filas enriquecidas : {enriq_filas}")
    for campo, n in sorted(campos_cnt.items(), key=lambda x: -x[1]):
        print(f"         {campo}: {n} filas completadas")

    df_vh    = df_maestro[df_maestro["RAMO"] == "VEHICULOS"].copy()
    df_otros = df_maestro[df_maestro["RAMO"] != "VEHICULOS"].copy()
    print(f"\n       VEHICULOS : {len(df_vh)} filas")
    print(f"       Otros ramos: {len(df_otros)} filas  {df_otros['RAMO'].unique().tolist()}")

    out_vh    = DIR_SALIDA / f"UNIFICADO_VEHICULOS_{fecha_str}.xlsx"
    out_otros = DIR_SALIDA / f"UNIFICADO_OTROS_{fecha_str}.xlsx"
    if not df_vh.empty:
        _guardar_vh(df_vh, out_vh)
        print(f"       Guardado: salida/{out_vh.name}")
    if not df_otros.empty:
        _guardar_otros(df_otros, out_otros)
        print(f"       Guardado: salida/{out_otros.name}")

    # [5/11] Generar CONTRATANTES
    print("\n[5/11] Generando CONTRATANTES...")
    df_cont, df_alertas_cont = generar_contratantes(df_maestro, CATALOGO_PATH, enr_d)
    bloq = (df_alertas_cont["NIVEL"] == "BLOQUEANTE").sum() if len(df_alertas_cont) else 0
    rev  = (df_alertas_cont["NIVEL"] == "REVISAR").sum() if len(df_alertas_cont) else 0
    info = (df_alertas_cont["NIVEL"] == "INFO").sum() if len(df_alertas_cont) else 0
    print(f"       Contratantes       : {len(df_cont)}")
    print(f"       Alertas BLOQUEANTE : {bloq}")
    print(f"       Alertas REVISAR    : {rev}")
    print(f"       Alertas INFO       : {info}")

    out_cont = DIR_SALIDA / f"CONTRATANTES_{fecha_str}.xlsx"
    with pd.ExcelWriter(out_cont, engine="openpyxl") as xw:
        df_cont.to_excel(xw, sheet_name="CONTRATANTES", index=False)
        df_alertas_cont.to_excel(xw, sheet_name="ALERTAS", index=False)

    ci_nivel = {}
    if len(df_alertas_cont):
        for _, _r in df_alertas_cont.iterrows():
            ci_nivel[str(_r["CI"]).strip()] = _r["NIVEL"]

    wb = load_workbook(out_cont)
    _aplicar_formato(wb)
    _formato_texto_cols(wb["CONTRATANTES"], ["CEDULA / RUC", "CELULAR", "CIUDAD", "CODIGO"])
    if ci_nivel:
        ws_cont = wb["CONTRATANTES"]
        hdr_c = [str(c.value or "") for c in ws_cont[1]]
        if "CEDULA / RUC" in hdr_c:
            ci_idx = hdr_c.index("CEDULA / RUC")
            for row in ws_cont.iter_rows(min_row=2):
                ci_val = str(row[ci_idx].value or "").strip()
                nivel = ci_nivel.get(ci_val)
                if nivel:
                    fill = (FILL_BLOQ if nivel == "BLOQUEANTE" else
                            FILL_REV  if nivel == "REVISAR"    else
                            FILL_ALERT)
                    for cell in row:
                        cell.fill = fill
    wb.save(out_cont)
    print(f"\n       Guardado: salida/{out_cont.name}")

    # [6/11] Generar POLIZAS VEHICULOS
    print("\n[6/11] Generando POLIZAS VEHICULOS (formato DBroker)...")
    out_pveh = None
    df_pol_vh = pd.DataFrame()
    pveh_alertas = []
    if not df_vh.empty:
        df_pol_vh, df_alert_pveh, pveh_alertas = generar_polizas_vh(df_vh, CATALOGO_PATH)
        print(f"       Polizas VH         : {len(df_pol_vh)}")
        print(f"       Alertas de mapeo   : {len(df_alert_pveh)}")

        out_pveh = DIR_SALIDA / f"POLIZAS_VH_{fecha_str}.xlsx"
        with pd.ExcelWriter(out_pveh, engine="openpyxl") as xw:
            df_pol_vh.to_excel(xw, sheet_name="VEHICULOS", index=False)
            df_alert_pveh.to_excel(xw, sheet_name="ALERTAS", index=False)

        wb = load_workbook(out_pveh)
        _aplicar_formato(wb)
        _formato_texto_cols(wb["VEHICULOS"], ["RUC_CED", "POLIZA", "FACTURA", "ANEXO"])
        _formato_decimal_cols(wb["VEHICULOS"], COLS_FINANCIERAS_PVEH)
        _formato_texto_cols(wb["VEHICULOS"], ["ANIO_FAB_VH"])
        ws_pveh = wb["VEHICULOS"]
        for row_idx, has_alert in enumerate(pveh_alertas, start=2):
            if has_alert:
                for cell in ws_pveh[row_idx]:
                    cell.fill = FILL_ALERT
        wb.save(out_pveh)
        print(f"\n       Guardado: salida/{out_pveh.name}")
    else:
        print("       Sin filas VEHICULOS, archivo no generado.")

    # [7/11] MATRICES VH
    print("\n[7/11] Generando MATRICES VH por aseguradora...")
    out_mtx = None
    if not df_vh.empty and out_pveh is not None:
        grupos_mtx, df_alert_mtx, cod_etiquetas = generar_matrices_vh(
            df_pol_vh, df_vh, CATALOGO_PATH, pveh_alertas
        )
        mtx_alertas = {}
        for _cod in list(grupos_mtx):
            _df_g = grupos_mtx[_cod]
            if "_ALERTA" in _df_g.columns:
                mtx_alertas[_cod] = _df_g["_ALERTA"].tolist()
                grupos_mtx[_cod] = _df_g.drop(columns=["_ALERTA"])
            else:
                mtx_alertas[_cod] = [False] * len(_df_g)

        print(f"       Hojas generadas    : {len(grupos_mtx)}")
        print(f"       Sin match          : {len(df_alert_mtx)}")
        for _cod, _df_g in sorted(grupos_mtx.items()):
            print(f"         {_cod}: {len(_df_g)} polizas")

        out_mtx = DIR_SALIDA / f"MATRICES_VH_{fecha_str}.xlsx"
        with pd.ExcelWriter(out_mtx, engine="openpyxl") as xw:
            for cod in sorted(grupos_mtx):
                etiq = cod_etiquetas.get(cod, "")
                nombre_hoja = f"{cod.replace('/', '-')} ({etiq})"[:31]
                grupos_mtx[cod].to_excel(xw, sheet_name=nombre_hoja, index=False)
            df_alert_mtx.to_excel(xw, sheet_name="ALERTAS", index=False)

        wb = load_workbook(out_mtx)
        _aplicar_formato(wb)
        for cod in sorted(grupos_mtx):
            etiq = cod_etiquetas.get(cod, "")
            nombre_hoja = f"{cod.replace('/', '-')} ({etiq})"[:31]
            ws = wb[nombre_hoja]
            _formato_texto_cols(ws, ["RUC_CED", "POLIZA", "FACTURA", "ANEXO", "ANIO_FAB_VH"])
            _formato_decimal_cols(ws, COLS_FINANCIERAS_PVEH)
            for row_idx, has_alert in enumerate(mtx_alertas.get(cod, []), start=2):
                if has_alert:
                    for cell in ws[row_idx]:
                        cell.fill = FILL_ALERT
        wb.save(out_mtx)
        print(f"\n       Guardado: salida/{out_mtx.name}")
    else:
        print("       Sin datos VEHICULOS, archivo no generado.")

    # [8/11] Generar POLIZAS INCENDIO/MULTIRIESGO
    print("\n[8/11] Generando POLIZAS INCENDIO/MULTIRIESGO (formato DBroker)...")
    out_pin = None
    df_pol_in = pd.DataFrame()
    pin_alertas = []
    if not df_otros.empty:
        df_pol_in, df_alert_in, pin_alertas = generar_polizas_in(df_otros, CATALOGO_PATH)
        if df_pol_in.empty:
            print("       Sin polizas INCENDIO/MULTIRIESGO en los archivos de entrada.")
        else:
            incendio = (df_pol_in["DESCRIPCION"] == "ESTRUCTURA").sum()
            mr       = (df_pol_in["DESCRIPCION"] == "ESTRUCTURA + CONTENIDO + ROBO").sum()
            print(f"       Polizas INCENDIO   : {incendio}")
            print(f"       Polizas MULTIRIESGO: {mr}")
            print(f"       Alertas de mapeo   : {len(df_alert_in)}")

            out_pin = DIR_SALIDA / f"POLIZAS_IN_{fecha_str}.xlsx"
            with pd.ExcelWriter(out_pin, engine="openpyxl") as xw:
                df_pol_in.to_excel(xw, sheet_name="INCENDIO MR", index=False)
                df_alert_in.to_excel(xw, sheet_name="ALERTAS", index=False)

            wb = load_workbook(out_pin)
            _aplicar_formato(wb)
            ws_in = wb["INCENDIO MR"]
            _formato_texto_cols(ws_in, ["RUC_CED", "POLIZA", "FACTURA", "ANEXO"])
            _formato_decimal_cols(ws_in, COLS_FINANCIERAS_PIN)
            for row_idx, has_alert in enumerate(pin_alertas, start=2):
                if has_alert:
                    for cell in ws_in[row_idx]:
                        cell.fill = FILL_ALERT
            wb.save(out_pin)
            print(f"\n       Guardado: salida/{out_pin.name}")
    else:
        print("       Sin datos OTROS, archivo no generado.")

    # [9/11] MATRICES INCENDIO/MULTIRIESGO
    print("\n[9/11] Generando MATRICES INCENDIO/MULTIRIESGO por aseguradora...")
    out_mtx_in = None
    if out_pin is not None and not df_pol_in.empty:
        _RAMOS_IN_MATX = {"INCENDIO Y ALIADAS", "MULTIRIESGO"}
        df_in_orig_matx = df_otros[
            df_otros["RAMO"].fillna("").str.strip().str.upper().isin(_RAMOS_IN_MATX)
        ].reset_index(drop=True)

        grupos_mtx_in, df_alert_mtx_in, cod_etiquetas_in = generar_matrices_in(
            df_pol_in, df_in_orig_matx, CATALOGO_PATH, pin_alertas
        )
        mtx_in_alertas = {}
        for _cod in list(grupos_mtx_in):
            _df_g = grupos_mtx_in[_cod]
            if "_ALERTA" in _df_g.columns:
                mtx_in_alertas[_cod] = _df_g["_ALERTA"].tolist()
                grupos_mtx_in[_cod] = _df_g.drop(columns=["_ALERTA"])
            else:
                mtx_in_alertas[_cod] = [False] * len(_df_g)

        print(f"       Hojas generadas    : {len(grupos_mtx_in)}")
        print(f"       Sin match          : {len(df_alert_mtx_in)}")
        for _cod, _df_g in sorted(grupos_mtx_in.items()):
            print(f"         {_cod}: {len(_df_g)} polizas")

        out_mtx_in = DIR_SALIDA / f"MATRICES_IN_{fecha_str}.xlsx"
        with pd.ExcelWriter(out_mtx_in, engine="openpyxl") as xw:
            for cod in sorted(grupos_mtx_in):
                etiq = cod_etiquetas_in.get(cod, "")
                nombre_hoja = f"{cod.replace('/', '-')} ({etiq})"[:31]
                grupos_mtx_in[cod].to_excel(xw, sheet_name=nombre_hoja, index=False)
            df_alert_mtx_in.to_excel(xw, sheet_name="ALERTAS", index=False)

        wb = load_workbook(out_mtx_in)
        _aplicar_formato(wb)
        for cod in sorted(grupos_mtx_in):
            etiq = cod_etiquetas_in.get(cod, "")
            nombre_hoja = f"{cod.replace('/', '-')} ({etiq})"[:31]
            ws = wb[nombre_hoja]
            _formato_texto_cols(ws, ["RUC_CED", "POLIZA", "FACTURA", "ANEXO"])
            _formato_decimal_cols(ws, COLS_FINANCIERAS_PIN)
            for row_idx, has_alert in enumerate(mtx_in_alertas.get(cod, []), start=2):
                if has_alert:
                    for cell in ws[row_idx]:
                        cell.fill = FILL_ALERT
        wb.save(out_mtx_in)
        print(f"\n       Guardado: salida/{out_mtx_in.name}")
    else:
        print("       Sin polizas INCENDIO/MR, archivo no generado.")

    # [10/11] Generar POLIZAS AP/VIDA
    print("\n[10/11] Generando POLIZAS AP/VIDA (formato DBroker)...")
    out_pav = None
    df_pol_ap          = pd.DataFrame()
    df_pol_vida        = pd.DataFrame()
    ap_alertas_flags   = []
    vida_alertas_flags = []
    if not df_otros.empty:
        df_pol_ap, df_pol_vida, df_alert_avida, ap_alertas_flags, vida_alertas_flags = generar_polizas_ap_vida(
            df_otros, CATALOGO_PATH
        )
        n_ap   = len(df_pol_ap)
        n_vida = len(df_pol_vida)

        if n_ap == 0 and n_vida == 0:
            print("       Sin polizas AP/VIDA en los archivos de entrada.")
        else:
            print(f"       Polizas AP         : {n_ap}")
            print(f"       Polizas VIDA       : {n_vida}")
            print(f"       Alertas de mapeo   : {len(df_alert_avida)}")

            out_pav = DIR_SALIDA / f"POLIZAS_AP_VIDA_{fecha_str}.xlsx"
            with pd.ExcelWriter(out_pav, engine="openpyxl") as xw:
                if n_ap > 0:
                    df_pol_ap.to_excel(xw, sheet_name="AP", index=False)
                if n_vida > 0:
                    df_pol_vida.to_excel(xw, sheet_name="VIDA", index=False)
                df_alert_avida.to_excel(xw, sheet_name="ALERTAS", index=False)

            wb = load_workbook(out_pav)
            _aplicar_formato(wb)
            if n_ap > 0:
                ws_ap = wb["AP"]
                _formato_texto_cols(ws_ap, ["RUC_CED", "POLIZA", "FACTURA", "ANEXO"])
                _formato_decimal_cols(ws_ap, COLS_FINANCIERAS_PIN)
                for row_idx, has_alert in enumerate(ap_alertas_flags, start=2):
                    if has_alert:
                        for cell in ws_ap[row_idx]:
                            cell.fill = FILL_ALERT
            if n_vida > 0:
                ws_vida = wb["VIDA"]
                _formato_texto_cols(ws_vida, ["RUC_CED", "POLIZA", "FACTURA", "ANEXO"])
                _formato_decimal_cols(ws_vida, COLS_FINANCIERAS_PIN)
                for row_idx, has_alert in enumerate(vida_alertas_flags, start=2):
                    if has_alert:
                        for cell in ws_vida[row_idx]:
                            cell.fill = FILL_ALERT
            if "ALERTAS" in wb.sheetnames:
                _colorear_alertas(wb["ALERTAS"])
            wb.save(out_pav)
            print(f"\n       Guardado: salida/{out_pav.name}")
    else:
        print("       Sin datos OTROS, archivo no generado.")

    # [11/11] MATRICES AP/VIDA
    print("\n[11/11] Generando MATRICES AP/VIDA por aseguradora...")
    out_mtx_av = None
    if out_pav is not None and (not df_pol_ap.empty or not df_pol_vida.empty):
        grupos_mtx_av, df_alert_mtx_av, cod_etiquetas_av = generar_matrices_ap_vida(
            df_pol_ap, df_pol_vida, df_otros, CATALOGO_PATH,
            ap_alertas_flags, vida_alertas_flags,
        )
        mtx_av_alertas = {}
        for _cod in list(grupos_mtx_av):
            _df_g = grupos_mtx_av[_cod]
            if "_ALERTA" in _df_g.columns:
                mtx_av_alertas[_cod] = _df_g["_ALERTA"].tolist()
                grupos_mtx_av[_cod] = _df_g.drop(columns=["_ALERTA"])
            else:
                mtx_av_alertas[_cod] = [False] * len(_df_g)

        print(f"       Hojas generadas    : {len(grupos_mtx_av)}")
        print(f"       Sin match          : {len(df_alert_mtx_av)}")
        for _cod, _df_g in sorted(grupos_mtx_av.items()):
            etiq_av = cod_etiquetas_av.get(_cod, "")
            print(f"         {_cod} ({etiq_av}): {len(_df_g)} polizas")

        out_mtx_av = DIR_SALIDA / f"MATRICES_AP_VIDA_{fecha_str}.xlsx"
        with pd.ExcelWriter(out_mtx_av, engine="openpyxl") as xw:
            for cod in sorted(grupos_mtx_av):
                etiq = cod_etiquetas_av.get(cod, "")
                nombre_hoja = f"{cod.replace('/', '-')} ({etiq})"[:31]
                grupos_mtx_av[cod].to_excel(xw, sheet_name=nombre_hoja, index=False)
            df_alert_mtx_av.to_excel(xw, sheet_name="ALERTAS", index=False)

        wb = load_workbook(out_mtx_av)
        _aplicar_formato(wb)
        for cod in sorted(grupos_mtx_av):
            etiq = cod_etiquetas_av.get(cod, "")
            nombre_hoja = f"{cod.replace('/', '-')} ({etiq})"[:31]
            ws = wb[nombre_hoja]
            _formato_texto_cols(ws, ["RUC_CED", "POLIZA", "FACTURA", "ANEXO"])
            _formato_decimal_cols(ws, COLS_FINANCIERAS_PIN)
            for row_idx, has_alert in enumerate(mtx_av_alertas.get(cod, []), start=2):
                if has_alert:
                    for cell in ws[row_idx]:
                        cell.fill = FILL_ALERT
        wb.save(out_mtx_av)
        print(f"\n       Guardado: salida/{out_mtx_av.name}")
    else:
        print("       Sin polizas AP/VIDA, archivo no generado.")

    # --- LOG DE PROCESO ---
    _RAMOS_IN_SET = {"INCENDIO Y ALIADAS", "MULTIRIESGO"}

    _ci_pveh = set()
    if out_pveh is not None:
        for _, _r in df_pol_vh.iterrows():
            _ci_pveh.add(str(_r["RUC_CED"]).strip())

    _ci_pin = set()
    if out_pin is not None and not df_pol_in.empty:
        for _, _r in df_pol_in.iterrows():
            _ci_pin.add(str(_r["RUC_CED"]).strip())

    _ci_pav = set()
    if out_pav is not None:
        for _, _r in df_pol_ap.iterrows():
            _ci_pav.add(str(_r["RUC_CED"]).strip())
        for _, _r in df_pol_vida.iterrows():
            _ci_pav.add(str(_r["RUC_CED"]).strip())

    _ci_cont = set(str(c).strip() for c in df_cont["CEDULA / RUC"])

    _log = []

    def _L(*args):
        _log.append(" ".join(str(a) for a in args))

    _L("=" * 70)
    _L(f"OPERACIONES DBroker  |  LOG  |  {hoy.strftime('%d/%m/%Y %H:%M')}")
    _L("=" * 70)
    _L()
    _L("[ ARCHIVOS DE ENTRADA ]")
    for _f in archivos_cash:
        _L(f"  {_f.name}")
    _L()
    _L("[ LECTURA DE ARCHIVOS CASH ]")
    for _linea in log_unif:
        _L(f"  {_linea}")
    _L(f"  Total filas en MAESTRO   : {len(df_maestro)}")
    _L(f"  CIs unicas               : {df_maestro['CI'].nunique()}")
    _L()
    _L("[ DISTRIBUCION POR RAMO ]")
    for _ramo, _cnt in df_maestro.groupby("RAMO", dropna=False).size().items():
        _L(f"  {str(_ramo or '(sin ramo)'):<28}  {_cnt} poliza(s)")
    _L()
    _L("[ DETALLE POR REGISTRO ]")
    _L(f"  {'CI':<16} {'POLIZA':<22} {'RAMO':<24} PVEH  PIN  PAV  CONT  ESTADO")
    _L("  " + "-" * 104)

    _filas_problema = []
    for _, _r in df_maestro.sort_values(["RAMO", "CI"]).iterrows():
        _ci      = str(_r["CI"] or "").strip()
        _poliza  = str(_r["POLIZA_EMITIDA"] or "(sin poliza)").strip()
        _ramo    = str(_r["RAMO"] or "").strip()
        _alertas = str(_r.get("ALERTAS") or "")
        _ramo_up = _ramo.upper()

        _pveh = "OK" if _ramo_up == "VEHICULOS" and _ci in _ci_pveh else (" - " if _ramo_up != "VEHICULOS" else "NO")
        _pin  = "OK" if _ramo_up in _RAMOS_IN_SET and _ci in _ci_pin else (" - " if _ramo_up not in _RAMOS_IN_SET else "NO")
        _pav  = "OK" if _ramo_up in ("AP", "VIDA") and _ci in _ci_pav else (" - " if _ramo_up not in ("AP", "VIDA") else "NO")
        _cont = "OK" if _ci in _ci_cont else "NO"

        if not _ramo:
            _estado = "RAMO VACIO - no procesado en POLIZAS"
            _filas_problema.append((_ci, _poliza, _ramo, _alertas, _estado))
        elif _ramo_up in ("AP", "VIDA"):
            if _pav == "NO":
                _estado = f"RAMO {_ramo}: no generado en POLIZAS_AP_VIDA"
                _filas_problema.append((_ci, _poliza, _ramo, _alertas, _estado))
            elif "BLOQUEANTE" in _alertas.upper():
                _estado = "CONTIENE ALERTA BLOQUEANTE"
            elif _alertas:
                _estado = "con alertas"
            else:
                _estado = "OK"
        elif _ramo_up not in ({"VEHICULOS"} | _RAMOS_IN_SET):
            _estado = f"RAMO '{_ramo}' desconocido - no procesado"
            _filas_problema.append((_ci, _poliza, _ramo, _alertas, _estado))
        elif "BLOQUEANTE" in _alertas.upper():
            _estado = "CONTIENE ALERTA BLOQUEANTE"
        elif _alertas:
            _estado = "con alertas"
        else:
            _estado = "OK"

        _al_corta = (_alertas[:55] + "...") if len(_alertas) > 55 else _alertas
        _L(f"  {_ci:<16} {_poliza:<22} {_ramo:<24} {_pveh:>4} {_pin:>4} {_pav:>4} {_cont:>4}  {_estado}")
        if _alertas:
            _L(f"  {'':>70}  >> {_al_corta}")
    _L()

    if _filas_problema:
        _L("[ REGISTROS NO PROCESADOS EN POLIZAS - REQUIEREN ATENCION ]")
        for _ci, _pol, _ramo, _al, _est in _filas_problema:
            _L(f"  CI={_ci}  POLIZA={_pol}  RAMO='{_ramo}'")
            _L(f"    Razon   : {_est}")
            if _al:
                _L(f"    Alertas : {_al[:120]}")
        _L()
    else:
        _L("[ REGISTROS NO PROCESADOS: ninguno - todos enrutados correctamente ]")
        _L()

    _L("[ RESUMEN DE TOTALES ]")
    _n_pveh = len(df_pol_vh) if out_pveh is not None else 0
    _n_pin  = len(df_pol_in) if out_pin is not None and not df_pol_in.empty else 0
    _n_pav  = (len(df_pol_ap) + len(df_pol_vida)) if out_pav is not None else 0
    _L(f"  Filas leidas del CASH        : {len(df_maestro)}")
    _L(f"  En POLIZAS VEHICULOS (PVEH)  : {_n_pveh}")
    _L(f"  En POLIZAS INCENDIO/MR (PIN) : {_n_pin}")
    _L(f"  En POLIZAS AP/VIDA (PAV)     : {_n_pav}")
    _L(f"  En CONTRATANTES              : {len(df_cont)}")
    _L(f"  Alertas BLOQUEANTE           : {bloq}")
    _L(f"  Alertas REVISAR              : {rev}")
    _L(f"  Alertas INFO                 : {info}")
    _n_prob = len(_filas_problema)
    if _n_prob:
        _L(f"  SIN RUTA A POLIZAS           : {_n_prob}  <-- REVISAR")
    _L()
    _L("[ ARCHIVOS GENERADOS ]")
    for _out in [out_vh, out_otros, out_cont, out_pveh, out_mtx, out_pin, out_mtx_in, out_pav, out_mtx_av]:
        if _out and _out.exists():
            _L(f"  salida/{_out.name}")
    _L("=" * 70)

    log_path = DIR_SALIDA / f"LOG_OPERACIONES_{fecha_str}.txt"
    with open(log_path, "w", encoding="utf-8") as _flog:
        _flog.write("\n".join(_log))
    print(f"\n       Log guardado: salida/{log_path.name}")

    print("\n" + "=" * 60)
    print("PROCESO OPERACIONES COMPLETADO")
    if not df_vh.empty:
        print(f"  salida/{out_vh.name}")
    if not df_otros.empty:
        print(f"  salida/{out_otros.name}")
    print(f"  salida/{out_cont.name}")
    if out_pveh:
        print(f"  salida/{out_pveh.name}")
    if out_mtx:
        print(f"  salida/{out_mtx.name}")
    if out_mtx_in:
        print(f"  salida/{out_mtx_in.name}")
    if out_pin:
        print(f"  salida/{out_pin.name}")
    if out_pav:
        print(f"  salida/{out_pav.name}")
    if out_mtx_av:
        print(f"  salida/{out_mtx_av.name}")
    print(f"  salida/{log_path.name}")
    if bloq > 0:
        print(f"\n  IMPORTANTE: {bloq} contratante(s) con alertas BLOQUEANTE.")
        print(f"  Abrir {out_cont.name}, hoja ALERTAS, y resolver antes de cargar.")
    if _n_prob > 0:
        print(f"\n  ATENCION: {_n_prob} poliza(s) sin ruta de salida. Ver LOG para detalle.")
    print("=" * 60)


if __name__ == "__main__":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    run()
