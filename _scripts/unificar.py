# -*- coding: utf-8 -*-
"""
Modulo: unificar.py
Lee archivos CASH de asesores (layout VH o IN) y devuelve un DataFrame maestro
con nombres de columna canonicos. No elimina ninguna fila: los problemas se
marcan en la columna ALERTAS.
"""
import re
import calendar
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
import pandas as pd

# ---------------------------------------------------------------
# Mapeo posicional: (nombre_canonico, indice_VH, indice_IN)
# None = columna no existe en ese layout
# ---------------------------------------------------------------
MAPEO = [
    # identificacion / cliente
    ("EJECUTIVO",              0,  0),
    ("CI",                     1,  1),
    ("GARANTIA",               2,  2),
    ("NOMBRE_CLIENTE",         3,  5),
    ("TIPO",                   4,  7),
    ("REGION",                 5,  None),
    ("CIUDAD",                 6,  6),
    ("OBS_POLIZA",             7,  8),
    ("CELULAR_1",              8,  9),
    ("CUENTA",                 9,  10),
    ("COMENTARIO",             10, 11),
    ("ESTADO_RENOVACION",      11, 12),
    # poliza emitida (valida para DBroker)
    ("POLIZA_EMITIDA",         23, 13),
    ("ASEGURADORA_EMISORA",    24, 14),
    ("RAMO",                   31, 21),
    ("VIGENCIA_DESDE",         33, 23),
    ("VIGENCIA_HASTA",         34, 24),
    ("CUOTA_INICIAL",          25, 15),
    ("FACTURA_ASEGURADORA",    26, 16),
    ("NUM_CUOTAS",             27, 17),
    ("FECHA_COBRO_CI",         28, 18),
    ("FECHA_COBRO_CF",         29, 19),
    # valores economicos
    ("VALOR_ASEGURADO",        12, 26),
    ("TASA",                   13, 29),
    ("PRIMA_NETA",             14, 31),
    ("COMISION",               15, 30),
    ("DERECHOS_EMISION",       16, 34),
    ("SUPER_BANCOS",           17, 32),
    ("SEGURO_CAMPESINO",       18, 33),
    ("IVA",                    19, 36),
    ("AXA",                    20, None),
    ("VIDA_ADICIONAL",         21, None),
    ("PRIMA_TOTAL",            22, 37),
    # poliza anterior (historico banco)
    ("POLIZA_ANTERIOR",        30, 20),
    ("ASEGURADORA_ANTERIOR",   32, 22),
    # analisis renovacion (solo VH)
    ("ULTIMO_VALOR_ASEG",      35, None),
    ("VALOR_DEPRECIADO",       36, None),
    ("TASA_PRESUPUESTO",       37, None),
    ("PN_PRESUPUESTO",         38, None),
    ("COM_PRESUPUESTO",        39, None),
    ("TASA_VIG_ANTERIOR",      40, None),
    ("TASA_RENOV_ASEG_ACTUAL", 41, None),
    # datos del vehiculo (solo VH)
    ("MARCA",                  42, None),
    ("MODELO",                 43, None),
    ("ANIO_VEHICULO",          44, None),
    ("MOTOR",                  45, None),
    ("CHASIS",                 46, None),
    ("COLOR",                  47, None),
    ("PLACA",                  48, None),
    # datos del inmueble / desgravamen (solo IN)
    ("DIRECCION_INMUEBLE",     None, 25),
    ("CONTENIDO",              None, 27),
    ("ROBO",                   None, 28),
    ("SUBTOTAL_IN",            None, 35),
    ("CUOTAS_IN",              None, 38),
    ("VALOR_CUOTA_IN",         None, 39),
    ("SALDO_IN",               None, 40),
    ("COMISION_DESGRAVAMEN",   None, 41),
    ("TASA_DESGRAVAMEN",       None, 42),
    ("PRIMA_NETA_DESGRAVAMEN", None, 43),
    ("SBS_DESGRAVAMEN",        None, 44),
    ("SC_DESGRAVAMEN",         None, 45),
    ("DE_DESGRAVAMEN",         None, 46),
    ("TOTAL_DESGRAVAMEN",      None, 47),
    ("CUOTAS_DESGRAVAMEN",     None, 48),
    ("VALOR_CUOTA_DESGRAVAMEN",None, 49),
    ("TOTAL_PRIMAS_A_PAGAR",   None, 50),
    ("TOTAL_DEBITO_MENSUAL",   None, 51),
    # credito bancario
    ("CUENTA_BANCO",           49, 59),
    ("NUM_PRESTAMO",           50, 52),
    ("CARTERA",                None, 53),
    ("ESTADO_CREDITO",         51, 54),
    ("FECHA_DESEMBOLSO",       52, 55),
    ("MONTO_CREDITO",          53, 56),
    ("SALDO_CREDITO",          54, 57),
    ("FECHA_VENC_CREDITO",     55, 58),
    # contacto / personales
    ("DIRECCION_DOM",          56, 60),
    ("TELF_DOM",               57, 61),
    ("DIRECCION_OFI",          58, 62),
    ("TELF_OFI",               59, 63),
    ("CELULAR_BANCO",          60, 64),
    ("CELULAR_2",              61, None),
    ("CORREO",                 62, 65),
    ("NOMBRE_OFICIAL",         63, 66),
    ("ASEGURADORA_BANCO",      64, 67),
    ("TELEFONOS_ADICIONALES",  65, 68),
    ("FECHA_NACIMIENTO",       66, 3),
    ("EDAD_ORIGINAL",          67, 4),
    ("GENERO",                 68, None),
    ("ESTADO_CIVIL",           69, None),
    ("PROFESION",              70, None),
]


def detectar_layout(headers):
    hu = [str(h).replace("\xa0", " ").strip().upper() if h else "" for h in headers]
    if len(hu) > 31 and hu[31] == "RAMO" and "MARCA" in hu:
        return "VH"
    if len(hu) > 21 and hu[21] == "RAMO" and "DIRECCION INMUEBLE" in hu:
        return "IN"
    return None


def normalizar_fecha(v, hoy):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None, "FECHA_NAC vacia"
    if isinstance(v, datetime):
        return v.date(), None
    if isinstance(v, date):
        return v, None
    if isinstance(v, (int, float)):
        try:
            d = from_excel(v)
            if isinstance(d, datetime):
                d = d.date()
            if 1900 < d.year <= hoy.year:
                return d, None
        except Exception:
            pass
        return None, f"FECHA_NAC serial invalido ({v})"
    s = str(v).strip()
    if "NO ENCONTRADO" in s.upper() or "NO VALIDO" in s.upper() or "NO VÁLIDO" in s.upper():
        return None, "FECHA_NAC no encontrada en fuente"
    for fmt in ("%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date(), None
        except ValueError:
            continue
    return None, f"FECHA_NAC formato no reconocido ('{s[:20]}')"


_RAMO_ALIASES = {
    "ACCIDENTES PERSONALES": "AP",
    "ACCIDENTE PERSONAL":    "AP",
    "ACCIDENTES":            "AP",
}


def _normalizar_ramo(v):
    """Normaliza el valor de RAMO del CASH a la forma canónica."""
    if not v:
        return None
    u = str(v).strip().upper()
    return _RAMO_ALIASES.get(u, str(v).strip()) or None


def _normalizar_pol_simple(s):
    """Normaliza una cadena de poliza simple (sin mezcla de ramos)."""
    s = re.sub(r"^(VH|AP|VD|IN|MR)\s+", "", s, flags=re.I)
    s = s.replace("/", "-")
    s = re.sub(r"\s+", "", s)
    tokens = [t for t in s.split("-") if t != ""]
    if not tokens:
        return None
    if len(tokens) == 1:
        return tokens[0]
    tiene_letra = any(re.search(r"[A-Za-z]", t) for t in tokens)
    if tiene_letra and len(tokens) == 2:
        return tokens[0] + tokens[1] + "-0"
    return "".join(tokens[:-1]) + "-" + tokens[-1]


_RAMO_LABEL = {
    "VEHICULOS":          ["VH"],
    "VIDA":               ["VIDA"],
    "AP":                 ["AP"],
    "INCENDIO Y ALIADAS": ["IN", "INCENDIO"],
    "MULTIRIESGO":        ["MR", "MULTIRIESGO"],
}


def _extraer_poliza_parte(s, ramo):
    """
    Cuando POLIZA_EMITIDA contiene polizas de varios ramos mezcladas
    (p.ej. 'VH 08D-4851161-VIDA 02D-0100060'), extrae solo la del RAMO actual.
    Retorna (poliza_str_o_None, alerta_o_None).
    """
    labels = _RAMO_LABEL.get(str(ramo or "").strip().upper(), [])
    # Separar el campo en pares (LABEL, VALOR) usando los labels como delimitadores
    # Ejemplo: ['', 'VH', '08D-4851161', 'VIDA', '02D-0100060 ']
    partes = re.split(r'[-\s]*(VH|VIDA|AP|IN|MR)\s+', s.strip(), flags=re.I)
    pares = {}
    for i in range(1, len(partes) - 1, 2):
        lbl = partes[i].strip().upper()
        val = partes[i + 1].strip().rstrip("-").strip()
        pares[lbl] = val
    for lbl in labels:
        if lbl in pares:
            return pares[lbl], None
    return None, f"POLIZA mixta sin parte reconocida para ramo '{ramo}'"


def normalizar_poliza(v, ramo=None):
    if v is None:
        return None, None
    s = str(v).strip()
    # Campo mixto: contiene labels de ramo como "VH", "VIDA", "AP", etc.
    if ramo and re.search(r'\b(VH|VIDA|AP|IN|MR)\b\s+', s, re.I):
        pol_part, al = _extraer_poliza_parte(s, ramo)
        if pol_part:
            norm = _normalizar_pol_simple(pol_part)
            return norm, al
        # No se pudo extraer — usar flujo normal y agregar alerta
        alerta_extra = al
    else:
        alerta_extra = None

    norm = _normalizar_pol_simple(s)
    if norm is None:
        return None, "POLIZA_EMITIDA vacia tras limpiar"
    alerta = alerta_extra
    # Detectar si quedaron tokens con letras de label embebidos (señal de mezcla no resuelta)
    if not alerta and re.search(r'\b(VIDA|VEHICULOS)\b', norm, re.I):
        alerta = "POLIZA con texto de ramo embebido — posible mezcla no resuelta"
    tokens = [t for t in norm.split("-") if t]
    if len(tokens) == 1 and not alerta:
        alerta = "POLIZA sin anexo - verificar"
    return norm, alerta


def _extraer_factura_parte(s, ramo):
    """
    Cuando el campo tiene varias facturas separadas por ';' (p.ej.
    'VH 001020000171915 ; VIDA 001025000223381'), extrae solo la
    correspondiente al RAMO actual.
    Retorna (digitos_str_o_None, alerta_o_None).
    """
    partes = [p.strip() for p in s.split(";")]
    labels = _RAMO_LABEL.get(str(ramo or "").strip().upper(), [])
    for lbl in labels:
        for parte in partes:
            if parte.upper().startswith(lbl):
                d = re.sub(r"\D", "", parte)
                return d, None
    # No match — devolver None para que el caller use el comportamiento normal
    return None, f"FACTURA mixta sin parte reconocida para ramo '{ramo}'"


def _padear_factura_sri(digitos):
    """
    Completa ceros faltantes en el secuencial SRI.
    Formato: 3 (establecimiento) + 3 (punto emision) + 9 (secuencial) = 15 digitos.
    Si hay entre 7 y 14 digitos, rellena la parte secuencial con ceros a la izquierda.
    """
    if len(digitos) == 15:
        return digitos
    if 7 <= len(digitos) < 15:
        return digitos[:6] + digitos[6:].zfill(9)
    return digitos   # < 7 o > 15: no se puede corregir automaticamente


def normalizar_factura(v, ramo=None, aseguradora=None):
    if v is None:
        return None, []
    s = str(v).strip()
    alertas = []
    if "PROFORMA" in s.upper():
        alertas.append("FACTURA: ES PROFORMA")

    es_alianza = "ALIANZA" in str(aseguradora or "").upper()

    # Campo con multiples facturas separadas por ';'
    if ";" in s and ramo:
        d, al = _extraer_factura_parte(s, ramo)
        if d:
            if es_alianza:
                d = _padear_factura_sri(d)
            if len(d) != 15:
                alertas.append(f"FACTURA con {len(d)} digitos (SRI exige 15)")
            return d, alertas
        if al:
            alertas.append(al)

    digitos = re.sub(r"\D", "", s)
    if not digitos:
        return None, alertas + ["FACTURA sin numero"]
    if es_alianza:
        digitos = _padear_factura_sri(digitos)
    if len(digitos) != 15:
        alertas.append(f"FACTURA con {len(digitos)} digitos (SRI exige 15)")
    return digitos, alertas


def _limpiar(v):
    if isinstance(v, str):
        v = v.replace("\xa0", " ").strip()
        return v if v else None
    return v


def unificar(archivos: list, hoy: date = None) -> tuple:
    """
    Lee una lista de rutas (Path o str) de archivos CASH y devuelve
    (df_maestro, log_list).
    """
    if hoy is None:
        hoy = date.today()

    registros, log = [], []

    for path in archivos:
        path = Path(path)
        nombre = path.name
        try:
            wb = load_workbook(path, data_only=True)
        except Exception as e:
            log.append(f"ERROR: no se pudo abrir '{nombre}': {e}")
            continue

        for ws in wb.worksheets:
            headers = [c.value for c in ws[1]]
            layout = detectar_layout(headers)
            if layout is None:
                log.append(f"AVISO: hoja '{ws.title}' de '{nombre}' no reconocida -> omitida")
                continue
            n = 0
            for fila_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if all(row[i] in (None, "") for i in (0, 1) if i < len(row)):
                    continue
                try:
                    reg = {
                        "ARCHIVO_ORIGEN": nombre,
                        "HOJA_ORIGEN": ws.title.strip(),
                        "FILA_ORIGEN": fila_idx,
                        "LAYOUT": layout,
                    }
                    for canon, ivh, iin in MAPEO:
                        idx = ivh if layout == "VH" else iin
                        reg[canon] = _limpiar(row[idx]) if idx is not None and idx < len(row) else None

                    reg["RAMO"] = _normalizar_ramo(reg.get("RAMO"))

                    # Fallback celular: col J vacía → usar col BM (CELULAR_BANCO)
                    if not reg.get("CELULAR_1") and reg.get("CELULAR_BANCO"):
                        reg["CELULAR_1"] = reg["CELULAR_BANCO"]

                    alertas = []

                    fnac, alerta_f = normalizar_fecha(reg["FECHA_NACIMIENTO"], hoy)
                    if alerta_f:
                        alertas.append(alerta_f)
                    reg["FECHA_NACIMIENTO_NORM"] = fnac
                    reg["EDAD_CALCULADA"] = (
                        hoy.year - fnac.year - ((hoy.month, hoy.day) < (fnac.month, fnac.day))
                        if fnac else None
                    )

                    ci = str(reg["CI"]).strip() if reg["CI"] is not None else ""
                    # Padding: 9 digitos -> CI ecuatoriana (10), 12 digitos -> RUC (13)
                    if ci.isdigit():
                        if len(ci) == 9:
                            ci = "0" + ci
                            alertas.append("CI completada: 9 digitos, se antepuso 0")
                        elif len(ci) == 12:
                            ci = "0" + ci
                            alertas.append("RUC completado: 12 digitos, se antepuso 0")
                    reg["CI"] = ci if ci else None
                    if not ci:
                        alertas.append("CI vacia")
                    elif not re.fullmatch(r"\d{10}|\d{13}", ci):
                        alertas.append(f"CI con formato invalido ({ci})")

                    # Normalizar TASA: quitar '%', convertir a float
                    # Si el valor esta en decimal (0.04) convertir a porcentaje (4.00)
                    tasa_raw = reg.get("TASA")
                    if tasa_raw is not None:
                        try:
                            tasa_clean = str(tasa_raw).replace("%", "").strip()
                            tasa_val = float(tasa_clean) if tasa_clean else None
                            if tasa_val is not None and 0 < tasa_val < 1:
                                tasa_val = round(tasa_val * 100, 4)  # 0.04 -> 4.0
                            reg["TASA"] = tasa_val
                        except (ValueError, TypeError):
                            alertas.append(f"TASA con formato no reconocido ({tasa_raw})")

                    obs = str(reg["OBS_POLIZA"] or "").upper()
                    if "ANULADA" in obs:
                        alertas.append("POLIZA ANULADA segun OBS_POLIZA (se conserva, revisar)")
                    if reg["RAMO"] is None:
                        # Para layout VH: inferir VEHICULOS si hay datos de marca/modelo
                        if layout == "VH" and reg.get("MARCA"):
                            reg["RAMO"] = "VEHICULOS"
                            alertas.append("RAMO vacio en CASH: asumido VEHICULOS por presencia de MARCA")
                        else:
                            alertas.append("RAMO vacio")
                    if reg["POLIZA_EMITIDA"] is None:
                        alertas.append("POLIZA_EMITIDA vacia")

                    reg["POLIZA_EMITIDA_ORIGINAL"] = reg["POLIZA_EMITIDA"]
                    pol_norm, alerta_pol = normalizar_poliza(
                        reg["POLIZA_EMITIDA"], ramo=reg.get("RAMO")
                    )
                    reg["POLIZA_EMITIDA"] = pol_norm
                    if alerta_pol:
                        alertas.append(alerta_pol)

                    reg["FACTURA_ASEGURADORA_ORIGINAL"] = reg["FACTURA_ASEGURADORA"]
                    fac_norm, alertas_fac = normalizar_factura(
                        reg["FACTURA_ASEGURADORA"], ramo=reg.get("RAMO"),
                        aseguradora=reg.get("ASEGURADORA_EMISORA"),
                    )
                    reg["FACTURA_ASEGURADORA"] = fac_norm
                    alertas.extend(alertas_fac)

                    # Homologacion de nombres de aseguradora
                    _aseg = str(reg["ASEGURADORA_EMISORA"] or "").strip().upper()
                    if _aseg in ("LATINA", "LATINA SEGUROS"):
                        reg["ASEGURADORA_EMISORA"] = "LATINA SEGUROS"

                    if reg["ASEGURADORA_EMISORA"] is None:
                        alertas.append("ASEGURADORA_EMISORA vacia")

                    # Validar vigencia: debe ser 365 o 366 dias
                    vd_raw = reg.get("VIGENCIA_DESDE")
                    vh_raw = reg.get("VIGENCIA_HASTA")
                    if vd_raw is not None and vh_raw is not None:
                        try:
                            def _to_date(v):
                                if isinstance(v, datetime):
                                    return v.date()
                                if isinstance(v, date):
                                    return v
                                return pd.to_datetime(str(v), dayfirst=True).date()
                            vd_d = _to_date(vd_raw)
                            vh_d = _to_date(vh_raw)
                            dias = (vh_d - vd_d).days
                            # Bisiesto: Feb 29 cae dentro del periodo
                            tiene_29feb = any(
                                calendar.isleap(y) and date(y, 2, 29) >= vd_d and date(y, 2, 29) < vh_d
                                for y in range(vd_d.year, vh_d.year + 1)
                                if calendar.isleap(y)
                            )
                            esperado = 366 if tiene_29feb else 365
                            if dias != esperado:
                                alertas.append(
                                    f"VIGENCIA: {dias} dias (se esperan {esperado})"
                                )
                        except Exception:
                            pass  # fechas no parseables ya generan otra alerta

                    g = (str(reg["GENERO"]) if reg["GENERO"] else "").upper()
                    if layout == "VH" and ("NO ENCONTRADO" in g or g == ""):
                        alertas.append("GENERO no disponible")

                    ec = (str(reg["ESTADO_CIVIL"]) if reg["ESTADO_CIVIL"] else "").upper()
                    if layout == "VH" and ("NO ENCONTRADO" in ec or ec == ""):
                        alertas.append("ESTADO_CIVIL no disponible")

                    if reg["DIRECCION_DOM"] is None:
                        alertas.append("DIRECCION_DOM vacia")

                    reg["ALERTAS"] = " | ".join(alertas) if alertas else None
                    reg["NUM_ALERTAS"] = len(alertas)
                    registros.append(reg)
                    n += 1
                except Exception as _e:
                    log.append(f"ERROR fila {fila_idx} de '{nombre}' [{ws.title}]: {type(_e).__name__}: {_e}")

            log.append(f"OK: '{nombre}' / '{ws.title}' -> layout {layout}, {n} filas")
        wb.close()

    df = pd.DataFrame(registros)
    if not df.empty:
        dup_ci = df["CI"].value_counts()
        df["CLIENTE_MULTIPOLIZA"] = df["CI"].map(
            lambda c: "SI" if c and dup_ci.get(c, 0) > 1 else "NO"
        )

        # Detectar POLIZA_EMITIDA duplicada entre fila VH y fila VIDA/AP del mismo CI
        _RAMOS_VIDA_AP = {"VIDA", "AP"}
        for ci, grupo in df.groupby("CI"):
            ramos_norm = grupo["RAMO"].fillna("").str.strip().str.upper()
            vh_mask    = ramos_norm == "VEHICULOS"
            vida_mask  = ramos_norm.isin(_RAMOS_VIDA_AP)
            if not (vh_mask.any() and vida_mask.any()):
                continue
            pols_vh = set(grupo.loc[vh_mask, "POLIZA_EMITIDA"].dropna())
            for idx in grupo[vida_mask].index:
                pol = df.at[idx, "POLIZA_EMITIDA"]
                if pol and pol in pols_vh:
                    nueva = (
                        "REVISAR: POLIZA_EMITIDA parece ser de VH "
                        f"({pol}) duplicada con fila VEHICULOS — verificar poliza VIDA/AP real"
                    )
                    existente = df.at[idx, "ALERTAS"]
                    existente = "" if (existente is None or pd.isna(existente)) else str(existente)
                    df.at[idx, "ALERTAS"] = (existente + " | " + nueva).strip(" | ")
                    prev_n = df.at[idx, "NUM_ALERTAS"]
                    df.at[idx, "NUM_ALERTAS"] = int(prev_n if (prev_n is not None and not pd.isna(prev_n)) else 0) + 1

    return df, log
